from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
original='slides-engine/fixtures/independent.pptx';out='validation/slides/package3-edited.pptx'
a,b=ZipFile(original),ZipFile(out)
assert set(a.namelist())==set(b.namelist())
changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
assert changed==['ppt/slides/slide1.xml'],changed
for n in b.namelist():
 if n.endswith('.xml') or n.endswith('.rels'):etree.fromstring(b.read(n))
p=Presentation(out);assert p.slides[0].shapes[0].text=='Somnia edited copy'
print('Independent python-pptx and XML roundtrip passed; only slide1.xml changed, every other uncompressed part byte-identical')
