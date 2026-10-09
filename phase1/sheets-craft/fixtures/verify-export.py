from pathlib import Path
from openpyxl import load_workbook
import sys
p=Path(sys.argv[1])
f=load_workbook(p); v=load_workbook(p,data_only=True)
assert f.sheetnames == ['Sales','Summary']
assert f['Sales']['B2'].value==7
assert f['Sales']['C2'].value=='=B2*10'
assert v['Sales']['C2'].value==70
assert v['Summary']['A1'].value==85
assert f['Summary']['B1'].value=='Grüße ☕'
assert f['Sales']['A1'].font.bold
assert f['Sales'].freeze_panes=='A2'
assert 'A5:C5' in f['Sales'].merged_cells
assert len(f['Sales']._charts)==1
print('PASS: independent openpyxl export checks; chart existence only, not visual fidelity.')
