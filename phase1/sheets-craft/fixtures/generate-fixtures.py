"""Independent XLSX producer, not GridCraft's own serializer. Requires openpyxl 3.1.5."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Reference
root = Path(__file__).resolve().parents[0]
w = Workbook(); s = w.active; s.title = 'Sales'
s.append(['Product', 'Quantity', 'Total'])
s.append(['Coffee', 2, '=B2*10']); s.append(['Tea', 3, '=B3*5'])
s['A1'].font = Font(bold=True); s.freeze_panes = 'A2'
s.merge_cells('A5:C5'); s['A5']='Independent openpyxl fixture'
chart = BarChart(); chart.add_data(Reference(s, min_col=2, min_row=1, max_row=3), titles_from_data=True); s.add_chart(chart,'E2')
t = w.create_sheet('Summary'); t['A1']='=SUM(Sales!C2:C3)'; t['B1']='Grüße ☕'
w.save(root/'fixture.xlsx')

# Package 2: formatting fixture (number formats, bold/italic, fills, alignment, widths, hidden column).
import datetime
from openpyxl.styles import Alignment, PatternFill
f = Workbook(); g = f.active; g.title = 'Formats'
g.append(['Label', 'Value', 'Note'])
g['A1'].font = Font(bold=True, italic=True, color='FFFF0000'); g['A1'].fill = PatternFill('solid', fgColor='FFFFF2CC')
g['A2'] = 'Percent'; g['B2'] = 0.256; g['B2'].number_format = '0.0%'
g['A3'] = 'Currency'; g['B3'] = 1234.5; g['B3'].number_format = '#,##0.00 "EUR"'
g['A4'] = 'Date'; g['B4'] = datetime.date(2026, 10, 9); g['B4'].number_format = 'yyyy-mm-dd'
g['A5'] = 'Negative red'; g['B5'] = -42; g['B5'].number_format = '0;[Red]-0'
g['A6'] = 'Centered'; g['B6'] = 'mid'; g['B6'].alignment = Alignment(horizontal='center')
g['A7'] = 'Right text'; g['B7'] = 'right'; g['B7'].alignment = Alignment(horizontal='right')
g['A8'] = 'Plain'; g['B8'] = 3.14159265358979
from openpyxl.styles import Border, Side
g['A9'] = 'Boxed'; g['A9'].border = Border(bottom=Side(style='medium', color='FF0000FF')); g['A9'].font = Font(name='Courier New', size=16)
g['A10'] = 'Merged across two columns'; g.merge_cells('A10:B10'); g['A10'].alignment = Alignment(horizontal='center')
for i in range(11, 61): g.cell(row=i, column=1, value=f'Row {i}')
g.freeze_panes = 'B2'
g.column_dimensions['A'].width = 30; g.column_dimensions['B'].width = 18; g.column_dimensions['C'].hidden = True
g.column_dimensions['D'].width = 8; g['D1'] = 'narrow'
f.save(root/'formats.xlsx')
